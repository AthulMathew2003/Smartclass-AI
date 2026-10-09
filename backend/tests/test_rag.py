import io
import uuid
import pytest
from unittest.mock import MagicMock, patch
from httpx import AsyncClient
import chromadb
import pypdf
import docx
import pptx
import openpyxl

from app.modules.users.service import UserService
from app.modules.organizations.service import OrganizationService
from app.modules.organizations.workspace_service import WorkspaceService
from app.modules.organizations.schemas import OnboardingRequest
from app.modules.subjects.service import SubjectService
from app.modules.subjects.schemas import SubjectCreateRequest
from app.modules.subjects.material_service import SubjectMaterialService
from app.modules.subjects.material_schemas import MaterialUploadUrlRequest, MaterialConfirmRequest
from app.modules.subjects.models import SubjectMaterialStatus, MaterialIndexStatus
from app.modules.rag.document_processor import DocumentProcessor, DocumentSection, compute_content_hash
from app.modules.rag.chunker import DocumentChunker, DocumentChunk, clean_text
from app.modules.rag.embeddings import EmbeddingService
from app.modules.rag.vector_store import VectorStoreService
from app.modules.rag.service import RagService
from app.modules.rag.exceptions import DocumentExtractionException, EmbeddingException


# ── Helpers to generate in-memory mock files ─────────────────────────

def create_sample_docx_bytes(paragraphs: list[str]) -> bytes:
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(1, 0).text = "Value A"
    table.cell(1, 1).text = "Value B"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_sample_pptx_bytes(slides_text: list[str]) -> bytes:
    prs = pptx.Presentation()
    blank_layout = prs.slide_layouts[6]
    for st in slides_text:
        slide = prs.slides.add_slide(blank_layout)
        tx_box = slide.shapes.add_textbox(100, 100, 200, 100)
        tf = tx_box.text_frame
        tf.text = st
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def create_sample_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DataSheet"
    ws.append(["ID", "Name", "Score"])
    ws.append(["1", "Alice", "95"])
    ws.append(["2", "Bob", "88"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 1. Document Extraction Tests ─────────────────────────────────────

def test_document_processor_text_and_csv():
    processor = DocumentProcessor()
    txt_bytes = b"Hello, world! This is a simple test document.\nWith multiple lines."
    sections = processor.extract_sections(txt_bytes, "text/plain", "notes.txt")
    assert len(sections) == 1
    assert "Hello, world!" in sections[0].text

    csv_bytes = b"ID,Subject,Teacher\n1,Mathematics,Prof. Euler\n2,Physics,Dr. Newton"
    csv_sections = processor.extract_sections(csv_bytes, "text/csv", "schedule.csv")
    assert len(csv_sections) == 1
    assert "Mathematics" in csv_sections[0].text


def test_document_processor_docx():
    processor = DocumentProcessor()
    docx_bytes = create_sample_docx_bytes(["Introduction to Database Systems", "Normalization eliminates redundancy."])
    sections = processor.extract_sections(docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "db.docx")
    assert len(sections) == 1
    assert "Introduction to Database Systems" in sections[0].text
    assert "Normalization eliminates redundancy." in sections[0].text
    assert "Header 1" in sections[0].text


def test_document_processor_pptx():
    processor = DocumentProcessor()
    pptx_bytes = create_sample_pptx_bytes(["Slide 1: Cloud Architecture", "Slide 2: Scalability and Elasticity"])
    sections = processor.extract_sections(pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation", "cloud.pptx")
    assert len(sections) == 2
    assert sections[0].slide_number == 1
    assert "Cloud Architecture" in sections[0].text
    assert sections[1].slide_number == 2
    assert "Scalability" in sections[1].text


def test_document_processor_xlsx():
    processor = DocumentProcessor()
    xlsx_bytes = create_sample_xlsx_bytes()
    sections = processor.extract_sections(xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "data.xlsx")
    assert len(sections) == 1
    assert "Alice" in sections[0].text
    assert "DataSheet" in sections[0].text


def test_document_processor_empty_and_unsupported():
    processor = DocumentProcessor()
    with pytest.raises(DocumentExtractionException):
        processor.extract_sections(b"", "text/plain", "empty.txt")

    with pytest.raises(DocumentExtractionException):
        processor.extract_sections(b"some content", "video/mp4", "video.mp4")


# ── 2. Text Chunking & Cleaning Tests ─────────────────────────────────

def test_clean_text():
    raw = "Hello   world!\r\n\r\n\r\nThis\tis\t\ta   test.\n\n\n\nDone."
    cleaned = clean_text(raw)
    assert "Hello world!" in cleaned
    assert "\r" not in cleaned
    assert "\t" not in cleaned
    assert "\n\n\n" not in cleaned


def test_document_chunker_deterministic():
    chunker = DocumentChunker(chunk_size=100, chunk_overlap=20)
    sections = [
        DocumentSection(text="A" * 150, page_number=1),
        DocumentSection(text="B" * 50, page_number=2),
    ]
    chunks = chunker.chunk_sections(sections)
    assert len(chunks) >= 2
    assert chunks[0].chunk_index == 0
    assert chunks[0].page_number == 1
    assert chunks[-1].page_number == 2


# ── 3. Mocked Gemini Embedding Service Tests ─────────────────────────

def test_embedding_service_mocked():
    mock_client = MagicMock()
    mock_resp = MagicMock()
    mock_item1 = MagicMock()
    mock_item1.values = [0.1, 0.2, 0.3]
    mock_resp.embeddings = [mock_item1]
    mock_client.models.embed_content.return_value = mock_resp

    service = EmbeddingService(api_key="test-gemini-key")
    service._client = mock_client

    # Query embedding
    q_vec = service.embed_query("What is indexing?")
    assert q_vec == [0.1, 0.2, 0.3]

    # Documents embedding
    mock_item2 = MagicMock()
    mock_item2.values = [0.4, 0.5, 0.6]
    mock_resp.embeddings = [mock_item1, mock_item2]
    docs_vec = service.embed_documents(["chunk 1", "chunk 2"])
    assert len(docs_vec) == 2


def test_embedding_service_missing_api_key():
    service = EmbeddingService(api_key="")
    with pytest.raises(EmbeddingException):
        service.embed_query("test query")


# ── 4. ChromaDB Vector Store Service Tests ───────────────────────────

def test_chromadb_vector_store():
    # Ephemeral in-memory Chroma client for test isolation
    client = chromadb.EphemeralClient()
    vector_store = VectorStoreService(client=client, collection_name="test_materials")

    org_id = uuid.uuid4()
    ws_id = uuid.uuid4()
    subj_id = uuid.uuid4()
    mat_id = uuid.uuid4()

    chunks = [
        DocumentChunk(chunk_index=0, text="Normalization is database design.", page_number=1),
        DocumentChunk(chunk_index=1, text="First normal form avoids repeating groups.", page_number=2),
    ]
    embeddings = [[0.1, 0.2, 0.3, 0.4], [0.5, 0.6, 0.7, 0.8]]

    count = vector_store.upsert_chunks(
        organization_id=org_id,
        workspace_id=ws_id,
        subject_id=subj_id,
        material_id=mat_id,
        material_title="Database Design",
        filename="dbms.pdf",
        content_hash="abc123hash",
        chunks=chunks,
        embeddings=embeddings,
    )
    assert count == 2

    # Search with correct tenant context
    results = vector_store.search(
        query_embedding=[0.1, 0.2, 0.3, 0.4],
        organization_id=org_id,
        workspace_id=ws_id,
        subject_id=subj_id,
        top_k=2,
    )
    assert len(results) == 2
    assert results[0]["material_title"] == "Database Design"
    assert results[0]["page_number"] in (1, 2)

    # Search with DIFFERENT subject (tenant isolation)
    wrong_subj_id = uuid.uuid4()
    isolated_results = vector_store.search(
        query_embedding=[0.1, 0.2, 0.3, 0.4],
        organization_id=org_id,
        workspace_id=ws_id,
        subject_id=wrong_subj_id,
        top_k=2,
    )
    assert len(isolated_results) == 0

    # Delete material vectors
    vector_store.delete_material_vectors(mat_id)
    post_delete_results = vector_store.search(
        query_embedding=[0.1, 0.2, 0.3, 0.4],
        organization_id=org_id,
        workspace_id=ws_id,
        subject_id=subj_id,
        top_k=2,
    )
    assert len(post_delete_results) == 0


# ── 5. Full End-to-End API Integration & Tenant Isolation Tests ──────

@pytest.mark.asyncio
async def test_rag_end_to_end_lifecycle_and_search(client: AsyncClient, db_session):
    # 1. Setup Owner user and organization
    user_service = UserService(db_session)
    user = await user_service.create_user(email="rag_owner@example.com", password="Password123!")
    await db_session.commit()

    login_res = await client.post(
        "/api/v1/auth/login",
        json={"email": "rag_owner@example.com", "password": "Password123!"}
    )
    token = login_res.json()["access_token"]

    org_service = OrganizationService(db_session)
    org = await org_service.complete_onboarding(
        user,
        OnboardingRequest(
            org_name="RAG Institute",
            org_slug="rag-inst",
            org_type="University",
            org_country="United States",
            org_state="California",
            org_city="Berkeley",
            org_timezone="America/Los_Angeles",
            workspace_name="CS Department",
            owner_first_name="Alan",
            owner_last_name="Turing"
        )
    )
    await db_session.commit()

    headers = {
        "Authorization": f"Bearer {token}",
        "X-Organization-Id": str(org.organization_id)
    }

    # 2. Get Workspace and Create Subject
    ws_res = await client.get("/api/v1/workspaces", headers=headers)
    ws_id = ws_res.json()[0]["workspace_id"]

    subj_res = await client.post(
        "/api/v1/subjects",
        headers=headers,
        json={"workspace_id": ws_id, "subject_name": "Artificial Intelligence"}
    )
    assert subj_res.status_code == 201
    subj_id = subj_res.json()["subject_id"]

    # 3. Request Upload URL for TXT material
    mock_storage = MagicMock()
    mock_storage.generate_upload_url.return_value = "https://s3.fake/upload"
    mock_storage.object_exists.return_value = True
    raw_material_content = b"Knowledge graphs and vector embeddings power modern retrieval augmented generation."
    mock_storage.get_object_bytes.return_value = raw_material_content

    # Mock Gemini Embeddings
    mock_embeddings = MagicMock(spec=EmbeddingService)
    mock_embeddings.embed_query.return_value = [0.1, 0.2, 0.3, 0.4]
    mock_embeddings.embed_documents.return_value = [[0.1, 0.2, 0.3, 0.4]]

    ephemeral_chroma = chromadb.EphemeralClient()
    test_vector_store = VectorStoreService(client=ephemeral_chroma, collection_name="e2e_rag")

    with patch("app.modules.subjects.material_service.S3StorageService", return_value=mock_storage), \
         patch("app.modules.rag.service.S3StorageService", return_value=mock_storage), \
         patch("app.modules.rag.service.EmbeddingService", return_value=mock_embeddings), \
         patch("app.modules.rag.service.VectorStoreService", return_value=test_vector_store), \
         patch("app.modules.rag.router.RagService") as MockRagRouterService:

        # Test upload url request
        upload_req = await client.post(
            f"/api/v1/subjects/{subj_id}/materials/upload-url",
            headers=headers,
            json={
                "title": "RAG Foundations",
                "category": "Notes",
                "original_filename": "rag_notes.txt",
                "content_type": "text/plain",
                "file_size": len(raw_material_content)
            }
        )
        assert upload_req.status_code == 201
        mat_id = upload_req.json()["material_id"]
        s3_key = upload_req.json()["s3_key"]

        # Confirm upload and trigger indexing
        confirm_req = await client.post(
            f"/api/v1/subjects/{subj_id}/materials/{mat_id}/confirm",
            headers=headers,
            json={"s3_key": s3_key}
        )
        assert confirm_req.status_code == 200
        mat_data = confirm_req.json()
        assert mat_data["index_status"] in ("processing", "indexed")

        # Direct search service test
        rag_service = RagService(
            db_session,
            storage=mock_storage,
            embeddings=mock_embeddings,
            vector_store=test_vector_store
        )
        await rag_service.index_material(
            org_id=org.organization_id,
            subject_id=uuid.UUID(subj_id),
            material_id=uuid.UUID(mat_id),
            requesting_user_id=user.user_id,
            is_org_admin=True,
            force=True
        )
        search_res = await rag_service.search(
            org_id=org.organization_id,
            subject_id=uuid.UUID(subj_id),
            query="retrieval augmented generation",
            requesting_user_id=user.user_id,
            is_org_admin=True,
            top_k=3
        )
        assert search_res.results_count == 1
        assert "retrieval augmented generation" in search_res.results[0].content
        assert search_res.results[0].material_title == "RAG Foundations"

        # Re-index material
        reindex_res = await rag_service.index_material(
            org_id=org.organization_id,
            subject_id=uuid.UUID(subj_id),
            material_id=uuid.UUID(mat_id),
            requesting_user_id=user.user_id,
            is_org_admin=True,
            force=True
        )
        assert reindex_res.material_index_status == MaterialIndexStatus.INDEXED

        # Status endpoint test
        status_res = await rag_service.get_index_status(
            org_id=org.organization_id,
            subject_id=uuid.UUID(subj_id),
            material_id=uuid.UUID(mat_id),
            requesting_user_id=user.user_id,
            is_org_admin=True
        )
        assert status_res.status == "indexed"
        assert status_res.chunk_count == 1


@pytest.mark.asyncio
async def test_tenant_isolation_cross_organization_blocked(client: AsyncClient, db_session):
    # Create Org 1 and User 1
    user_service = UserService(db_session)
    user1 = await user_service.create_user(email="org1_user@example.com", password="Password123!")
    user2 = await user_service.create_user(email="org2_user@example.com", password="Password123!")
    await db_session.commit()

    org_service = OrganizationService(db_session)
    org1 = await org_service.complete_onboarding(
        user1,
        OnboardingRequest(
            org_name="Org One",
            org_slug="org-one",
            org_type="School",
            org_country="US",
            org_state="NY",
            org_city="NYC",
            org_timezone="America/New_York",
            workspace_name="WS 1",
            owner_first_name="User",
            owner_last_name="One"
        )
    )
    org2 = await org_service.complete_onboarding(
        user2,
        OnboardingRequest(
            org_name="Org Two",
            org_slug="org-two",
            org_type="School",
            org_country="US",
            org_state="CA",
            org_city="SF",
            org_timezone="America/Los_Angeles",
            workspace_name="WS 2",
            owner_first_name="User",
            owner_last_name="Two"
        )
    )
    await db_session.commit()

    # Create Subject in Org 1
    subj_service = SubjectService(db_session)
    ws_service = WorkspaceService(db_session)
    workspaces = await ws_service.list_workspaces(org1.organization_id)
    ws1 = workspaces[0]
    subj1 = await subj_service.create_subject(
        org1.organization_id,
        ws1.workspace_id,
        user1.user_id,
        SubjectCreateRequest(
            workspace_id=ws1.workspace_id,
            subject_name="Biology 101"
        )
    )
    await db_session.commit()

    # User 2 tries to search Subject 1 using Org 2 headers
    login2 = await client.post("/api/v1/auth/login", json={"email": "org2_user@example.com", "password": "Password123!"})
    token2 = login2.json()["access_token"]
    headers2 = {
        "Authorization": f"Bearer {token2}",
        "X-Organization-Id": str(org2.organization_id)
    }

    mock_embeddings = MagicMock(spec=EmbeddingService)
    mock_embeddings.embed_query.return_value = [0.1, 0.2, 0.3, 0.4]

    with patch("app.modules.rag.service.EmbeddingService", return_value=mock_embeddings):
        res = await client.post(
            "/api/v1/rag/search",
            headers=headers2,
            json={"subject_id": str(subj1.subject_id), "query": "Cell division"}
        )
        # Must be 404 or 403 (tenant isolation enforced)
        assert res.status_code in (403, 404)
