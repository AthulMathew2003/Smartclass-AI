import io
import csv
import hashlib
from dataclasses import dataclass
from typing import List, Optional
import pypdf
import docx
import pptx
import openpyxl

from app.modules.rag.exceptions import DocumentExtractionException


@dataclass
class DocumentSection:
    text: str
    page_number: Optional[int] = None
    slide_number: Optional[int] = None


def compute_content_hash(file_bytes: bytes) -> str:
    """Calculate SHA-256 hash of raw file bytes."""
    return hashlib.sha256(file_bytes).hexdigest()


class DocumentProcessor:
    """
    Service responsible for extracting text from supported file formats (PDF, DOCX, PPTX, TXT, CSV, XLSX)
    while preserving page/slide boundaries for downstream citation.
    """

    def extract_sections(self, file_bytes: bytes, mime_type: str, filename: str) -> List[DocumentSection]:
        if not file_bytes:
            raise DocumentExtractionException("File is empty (0 bytes).")

        clean_mime = (mime_type or "").strip().lower()
        lower_name = (filename or "").strip().lower()

        try:
            if clean_mime == "application/pdf" or lower_name.endswith(".pdf"):
                return self._extract_pdf(file_bytes)
            elif (
                clean_mime in (
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    "application/msword"
                )
                or lower_name.endswith((".docx", ".doc"))
            ):
                return self._extract_docx(file_bytes)
            elif (
                clean_mime in (
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    "application/vnd.ms-powerpoint"
                )
                or lower_name.endswith((".pptx", ".ppt"))
            ):
                return self._extract_pptx(file_bytes)
            elif clean_mime == "text/csv" or lower_name.endswith(".csv"):
                return self._extract_csv(file_bytes)
            elif (
                clean_mime in (
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "application/vnd.ms-excel"
                )
                or lower_name.endswith((".xlsx", ".xls"))
            ):
                return self._extract_xlsx(file_bytes)
            elif clean_mime == "text/plain" or lower_name.endswith(".txt"):
                return self._extract_txt(file_bytes)
            else:
                raise DocumentExtractionException(
                    f"Unsupported MIME type '{mime_type}' for text extraction."
                )
        except DocumentExtractionException:
            raise
        except Exception as e:
            raise DocumentExtractionException(f"Failed to extract text from document: {str(e)}")

    def _extract_pdf(self, file_bytes: bytes) -> List[DocumentSection]:
        sections: List[DocumentSection] = []
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        
        if len(reader.pages) == 0:
            raise DocumentExtractionException("PDF contains no pages.")

        for idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                sections.append(DocumentSection(text=text, page_number=idx))

        if not sections or not any(s.text.strip() for s in sections):
            raise DocumentExtractionException(
                "PDF contains no extractable text (it may be a scanned image or protected document)."
            )

        return sections

    def _extract_docx(self, file_bytes: bytes) -> List[DocumentSection]:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
        except Exception as e:
            raise DocumentExtractionException(f"Invalid or corrupted Word document: {str(e)}")

        paragraphs_text: List[str] = []
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                paragraphs_text.append(p.text.strip())

        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs_text.append(row_text)

        full_text = "\n\n".join(paragraphs_text).strip()
        if not full_text:
            raise DocumentExtractionException("Word document contains no extractable text.")

        return [DocumentSection(text=full_text, page_number=None)]

    def _extract_pptx(self, file_bytes: bytes) -> List[DocumentSection]:
        try:
            prs = pptx.Presentation(io.BytesIO(file_bytes))
        except Exception as e:
            raise DocumentExtractionException(f"Invalid or corrupted PowerPoint presentation: {str(e)}")

        sections: List[DocumentSection] = []
        for idx, slide in enumerate(prs.slides, start=1):
            slide_texts: List[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)
            if slide_texts:
                combined_text = "\n".join(slide_texts).strip()
                if combined_text:
                    sections.append(DocumentSection(text=combined_text, slide_number=idx))

        if not sections:
            raise DocumentExtractionException("PowerPoint presentation contains no extractable text.")

        return sections

    def _extract_txt(self, file_bytes: bytes) -> List[DocumentSection]:
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("latin-1")
            except Exception as e:
                raise DocumentExtractionException(f"Unable to decode text file: {str(e)}")

        clean_text = text.strip()
        if not clean_text:
            raise DocumentExtractionException("Text document is empty.")

        return [DocumentSection(text=clean_text, page_number=None)]

    def _extract_csv(self, file_bytes: bytes) -> List[DocumentSection]:
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1")

        lines: List[str] = []
        reader = csv.reader(io.StringIO(text))
        for row in reader:
            row_str = " | ".join(c.strip() for c in row if c.strip())
            if row_str:
                lines.append(row_str)

        full_text = "\n".join(lines).strip()
        if not full_text:
            raise DocumentExtractionException("CSV file contains no extractable text.")

        return [DocumentSection(text=full_text, page_number=None)]

    def _extract_xlsx(self, file_bytes: bytes) -> List[DocumentSection]:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception as e:
            raise DocumentExtractionException(f"Invalid or corrupted Excel spreadsheet: {str(e)}")

        sections: List[DocumentSection] = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            sheet_lines: List[str] = []
            for row in sheet.iter_rows(values_only=True):
                row_vals = [str(val).strip() for val in row if val is not None and str(val).strip()]
                if row_vals:
                    sheet_lines.append(" | ".join(row_vals))

            if sheet_lines:
                sheet_text = f"Sheet: {sheet_name}\n" + "\n".join(sheet_lines)
                sections.append(DocumentSection(text=sheet_text, page_number=None))

        if not sections:
            raise DocumentExtractionException("Excel document contains no extractable cell data.")

        return sections
