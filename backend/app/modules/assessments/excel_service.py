import io
import re
import uuid
from decimal import Decimal, InvalidOperation
from typing import List, Dict, Any, Tuple, Optional, Set
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.modules.assessments.models import QuestionType
from app.modules.assessments.schemas import (
    QuestionOptionCreate,
    validate_question_options_rules,
    BulkUploadValidationErrorItem,
    BulkUploadPreviewItem,
)
from app.core.exceptions import ValidationException


MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_QUESTIONS_LIMIT = 500

SUPPORTED_TYPE_MAP = {
    "mcq_single": QuestionType.MCQ_SINGLE,
    "mcq": QuestionType.MCQ_SINGLE,
    "single": QuestionType.MCQ_SINGLE,
    "single_choice": QuestionType.MCQ_SINGLE,
    "single-choice": QuestionType.MCQ_SINGLE,
    "mcq_multiple": QuestionType.MCQ_MULTIPLE,
    "multiple": QuestionType.MCQ_MULTIPLE,
    "multiple_choice": QuestionType.MCQ_MULTIPLE,
    "multiple-choice": QuestionType.MCQ_MULTIPLE,
    "multi": QuestionType.MCQ_MULTIPLE,
    "multi_select": QuestionType.MCQ_MULTIPLE,
    "true_false": QuestionType.TRUE_FALSE,
    "tf": QuestionType.TRUE_FALSE,
    "true/false": QuestionType.TRUE_FALSE,
    "boolean": QuestionType.TRUE_FALSE,
    "short_answer": QuestionType.SHORT_ANSWER,
    "short": QuestionType.SHORT_ANSWER,
    "short-answer": QuestionType.SHORT_ANSWER,
    "sa": QuestionType.SHORT_ANSWER,
    "text": QuestionType.SHORT_ANSWER,
}


class BulkQuestionImportService:
    @staticmethod
    def generate_excel_template() -> bytes:
        """
        Generate a standardized, beautifully styled Excel template with:
        - Sheet 1: 'Questions' with column headers and sample rows.
        - Sheet 2: 'Instructions' with full rules and format documentation.
        """
        wb = openpyxl.Workbook()

        # ── Sheet 1: Questions ──────────────────────────────────
        ws_questions = wb.active
        ws_questions.title = "Questions"

        headers = [
            "question",
            "question_type",
            "marks",
            "option_a",
            "option_b",
            "option_c",
            "option_d",
            "correct_answer",
            "explanation",
            "order"
        ]

        # Header styling
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # Slate 800
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        ws_questions.append(headers)
        ws_questions.row_dimensions[1].height = 28

        for col_idx in range(1, len(headers) + 1):
            cell = ws_questions.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        # Example rows
        sample_rows = [
            [
                "What is the capital of India?",
                "mcq_single",
                2.0,
                "New Delhi",
                "Mumbai",
                "Kolkata",
                "Chennai",
                "A",
                "New Delhi is the official capital of India.",
                1
            ],
            [
                "Which of the following are compiled programming languages?",
                "mcq_multiple",
                3.0,
                "C++",
                "Python",
                "Rust",
                "HTML",
                "A, C",
                "C++ and Rust are compiled languages. Python is interpreted and HTML is a markup language.",
                2
            ],
            [
                "In Python, tuples are immutable data structures.",
                "true_false",
                1.0,
                "",
                "",
                "",
                "",
                "TRUE",
                "Tuples cannot be modified after creation in Python.",
                3
            ],
            [
                "Explain the primary difference between a process and a thread.",
                "short_answer",
                5.0,
                "",
                "",
                "",
                "",
                "",
                "A process has its own memory address space, whereas threads share the memory space of their parent process.",
                4
            ]
        ]

        sample_font = Font(name="Calibri", size=10)
        for row_data in sample_rows:
            ws_questions.append(row_data)
            row_idx = ws_questions.max_row
            ws_questions.row_dimensions[row_idx].height = 24
            for col_idx, _ in enumerate(row_data, start=1):
                cell = ws_questions.cell(row=row_idx, column=col_idx)
                cell.font = sample_font
                cell.border = thin_border
                cell.alignment = left_align

        # Auto-adjust column widths
        col_widths = {
            "A": 45,  # question
            "B": 18,  # question_type
            "C": 10,  # marks
            "D": 22,  # option_a
            "E": 22,  # option_b
            "F": 22,  # option_c
            "G": 22,  # option_d
            "H": 18,  # correct_answer
            "I": 35,  # explanation
            "J": 10   # order
        }
        for col_letter, width in col_widths.items():
            ws_questions.column_dimensions[col_letter].width = width

        # ── Sheet 2: Instructions ───────────────────────────────
        ws_instructions = wb.create_sheet(title="Instructions")

        title_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")  # Blue 500
        title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")

        ws_instructions.merge_cells("A1:D1")
        title_cell = ws_instructions["A1"]
        title_cell.value = "SmartClass AI — Question Bulk Upload Guidelines"
        title_cell.font = title_font
        title_cell.fill = title_fill
        title_cell.alignment = center_align
        ws_instructions.row_dimensions[1].height = 36

        instruction_sections = [
            ("Column Name", "Required", "Supported Types / Format", "Description & Rules"),
            ("question", "YES", "Text (Max 5,000 chars)", "The main question prompt. Whitespace is automatically trimmed."),
            ("question_type", "YES", "mcq_single | mcq_multiple | true_false | short_answer", "Supported types: 'mcq_single' (Single choice), 'mcq_multiple' (Multiple choice), 'true_false' (True/False), 'short_answer' (Short answer/essay)."),
            ("marks", "YES", "Positive Number (e.g. 1, 2.5, 5)", "Marks awarded for this question. Must be greater than 0."),
            ("option_a", "Conditional", "Text", "Option A text. Required for mcq_single and mcq_multiple."),
            ("option_b", "Conditional", "Text", "Option B text. Required for mcq_single and mcq_multiple."),
            ("option_c", "Optional", "Text", "Option C text. Optional additional choice for MCQ questions."),
            ("option_d", "Optional", "Text", "Option D text. Optional additional choice for MCQ questions."),
            ("correct_answer", "Conditional", "Keys (A, B, C, D) or TRUE / FALSE", "For mcq_single: exactly 1 key (e.g. 'A'). For mcq_multiple: comma-separated keys (e.g. 'A, C'). For true_false: 'TRUE' or 'FALSE'. For short_answer: leave blank."),
            ("explanation", "Optional", "Text (Max 5,000 chars)", "Teacher-only explanation displayed during review or result view."),
            ("order", "Optional", "Integer >= 1", "Optional custom display order. If left blank, rows follow the spreadsheet order.")
        ]

        table_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        table_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for r_idx, row_vals in enumerate(instruction_sections, start=3):
            ws_instructions.append(list(row_vals))
            ws_instructions.row_dimensions[r_idx].height = 24
            for c_idx in range(1, len(row_vals) + 1):
                cell = ws_instructions.cell(row=r_idx, column=c_idx)
                cell.border = thin_border
                if r_idx == 3:
                    cell.fill = table_header_fill
                    cell.font = table_header_font
                    cell.alignment = center_align
                else:
                    cell.font = sample_font
                    cell.alignment = left_align

        ws_instructions.column_dimensions["A"].width = 18
        ws_instructions.column_dimensions["B"].width = 14
        ws_instructions.column_dimensions["C"].width = 35
        ws_instructions.column_dimensions["D"].width = 65

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def parse_and_validate_workbook(
        file_bytes: bytes,
        existing_assessment_question_texts: Optional[Set[str]] = None
    ) -> Tuple[List[Dict[str, Any]], List[BulkUploadValidationErrorItem], Decimal]:
        """
        Safely parse and validate Excel file content.
        Returns:
            - parsed_questions: list of validated question dicts ready for insertion
            - errors: list of row-level error objects
            - total_marks: sum of all marks in the validated questions
        """
        if not file_bytes:
            raise ValidationException("Uploaded file is empty.")

        if len(file_bytes) > MAX_FILE_SIZE_BYTES:
            raise ValidationException(f"File size exceeds maximum limit of {MAX_FILE_SIZE_BYTES // (1024 * 1024)} MB.")

        try:
            # data_only=True ensures no formula execution / values only
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        except Exception as e:
            raise ValidationException(f"Invalid Excel (.xlsx) file. Could not read workbook: {str(e)}")

        # Find the Questions sheet
        sheet_name = None
        for name in wb.sheetnames:
            if name.strip().lower() == "questions":
                sheet_name = name
                break
        if not sheet_name:
            # Fallback to the first sheet if only 1 sheet or no sheet explicitly named Questions
            sheet_name = wb.sheetnames[0]

        ws = wb[sheet_name]

        # Read header row
        rows_iter = ws.iter_rows(values_only=True)
        try:
            raw_headers = next(rows_iter)
        except StopIteration:
            raise ValidationException("The spreadsheet is empty.")

        if not raw_headers or all(h is None for h in raw_headers):
            raise ValidationException("The spreadsheet header row is empty.")

        # Normalize headers
        header_map: Dict[str, int] = {}
        for idx, col in enumerate(raw_headers):
            if col is not None:
                norm = str(col).strip().lower().replace(" ", "_").replace("-", "_")
                header_map[norm] = idx

        # Header aliases mapping
        col_aliases = {
            "question": ["question", "question_text", "prompt", "text"],
            "question_type": ["question_type", "type", "q_type"],
            "marks": ["marks", "default_marks", "score", "points", "question_marks"],
            "option_a": ["option_a", "option1", "option_1", "opt_a", "a"],
            "option_b": ["option_b", "option2", "option_2", "opt_b", "b"],
            "option_c": ["option_c", "option3", "option_3", "opt_c", "c"],
            "option_d": ["option_d", "option4", "option_4", "opt_d", "d"],
            "option_e": ["option_e", "option5", "option_5", "opt_e", "e"],
            "option_f": ["option_f", "option6", "option_6", "opt_f", "f"],
            "correct_answer": ["correct_answer", "answer", "correct", "correct_option", "correct_answers", "key"],
            "explanation": ["explanation", "question_explanation", "solution", "rationale"],
            "order": ["order", "question_order", "seq", "sequence", "index"]
        }

        resolved_indices: Dict[str, Optional[int]] = {}
        for std_key, aliases in col_aliases.items():
            resolved_idx = None
            for alias in aliases:
                if alias in header_map:
                    resolved_idx = header_map[alias]
                    break
            resolved_indices[std_key] = resolved_idx

        # Validate mandatory columns
        required_columns = ["question", "question_type", "marks"]
        for req in required_columns:
            if resolved_indices[req] is None:
                raise ValidationException(f"Missing required column: '{req}'. Please verify your template headers.")

        parsed_questions: List[Dict[str, Any]] = []
        errors: List[BulkUploadValidationErrorItem] = []
        total_marks = Decimal("0.00")

        seen_in_sheet_questions: Dict[str, int] = {}
        existing_texts = existing_assessment_question_texts or set()

        row_num = 1  # header was row 1
        for row_values in rows_iter:
            row_num += 1
            if not row_values or all(v is None or str(v).strip() == "" for v in row_values):
                continue  # skip empty row

            if len(parsed_questions) >= MAX_QUESTIONS_LIMIT:
                errors.append(BulkUploadValidationErrorItem(
                    row=row_num,
                    field="file",
                    message=f"Workbook exceeds maximum limit of {MAX_QUESTIONS_LIMIT} questions per upload."
                ))
                break

            def get_cell(std_key: str) -> Optional[Any]:
                idx = resolved_indices.get(std_key)
                if idx is not None and idx < len(row_values):
                    val = row_values[idx]
                    if val is not None:
                        val_str = str(val).strip()
                        return val_str if val_str != "" else None
                return None

            row_error_before = len(errors)

            # 1. Question Text
            q_text = get_cell("question")
            if not q_text:
                errors.append(BulkUploadValidationErrorItem(
                    row=row_num,
                    field="question",
                    message="Question text cannot be empty."
                ))
            elif len(q_text) > 5000:
                errors.append(BulkUploadValidationErrorItem(
                    row=row_num,
                    field="question",
                    message="Question text exceeds maximum length of 5000 characters."
                ))

            # Duplicate detection within spreadsheet
            if q_text:
                norm_q = q_text.lower()
                if norm_q in seen_in_sheet_questions:
                    prev_row = seen_in_sheet_questions[norm_q]
                    errors.append(BulkUploadValidationErrorItem(
                        row=row_num,
                        field="question",
                        message=f"Duplicate question detected in row {row_num} (matches row {prev_row})."
                    ))
                else:
                    seen_in_sheet_questions[norm_q] = row_num

                if norm_q in existing_texts:
                    errors.append(BulkUploadValidationErrorItem(
                        row=row_num,
                        field="question",
                        message="Question already exists in this assessment."
                    ))

            # 2. Question Type
            raw_type = get_cell("question_type")
            q_type: Optional[QuestionType] = None
            if not raw_type:
                errors.append(BulkUploadValidationErrorItem(
                    row=row_num,
                    field="question_type",
                    message="Question type cannot be empty."
                ))
            else:
                norm_type = raw_type.lower().replace(" ", "_").replace("-", "_")
                if norm_type in SUPPORTED_TYPE_MAP:
                    q_type = SUPPORTED_TYPE_MAP[norm_type]
                else:
                    errors.append(BulkUploadValidationErrorItem(
                        row=row_num,
                        field="question_type",
                        message=f"Unsupported question type '{raw_type}'. Supported types: mcq_single, mcq_multiple, true_false, short_answer."
                    ))

            # 3. Marks
            raw_marks = get_cell("marks")
            marks_val: Optional[Decimal] = None
            if not raw_marks:
                errors.append(BulkUploadValidationErrorItem(
                    row=row_num,
                    field="marks",
                    message="Marks cannot be empty."
                ))
            else:
                try:
                    marks_val = Decimal(str(raw_marks))
                    if marks_val <= Decimal("0.00"):
                        errors.append(BulkUploadValidationErrorItem(
                            row=row_num,
                            field="marks",
                            message="Marks must be greater than 0."
                        ))
                    elif marks_val > Decimal("99999.99"):
                        errors.append(BulkUploadValidationErrorItem(
                            row=row_num,
                            field="marks",
                            message="Marks cannot exceed 99999.99."
                        ))
                    else:
                        marks_val = marks_val.quantize(Decimal("0.01"))
                except (InvalidOperation, ValueError):
                    errors.append(BulkUploadValidationErrorItem(
                        row=row_num,
                        field="marks",
                        message=f"Invalid marks value '{raw_marks}'. Must be a valid positive number."
                    ))

            # 4. Explanation & Order
            explanation = get_cell("explanation")
            raw_order = get_cell("order")
            order_val: Optional[int] = None
            if raw_order:
                try:
                    order_val = int(float(raw_order))
                    if order_val < 1:
                        errors.append(BulkUploadValidationErrorItem(
                            row=row_num,
                            field="order",
                            message="Question order must be an integer >= 1."
                        ))
                except ValueError:
                    errors.append(BulkUploadValidationErrorItem(
                        row=row_num,
                        field="order",
                        message=f"Invalid order value '{raw_order}'. Must be an integer."
                    ))

            # 5. Options & Correct Answer by Question Type
            options_list: List[QuestionOptionCreate] = []
            correct_ans_raw = get_cell("correct_answer")

            if q_type == QuestionType.MCQ_SINGLE:
                opt_a = get_cell("option_a")
                opt_b = get_cell("option_b")
                opt_c = get_cell("option_c")
                opt_d = get_cell("option_d")
                opt_e = get_cell("option_e")
                opt_f = get_cell("option_f")

                raw_opts = [
                    ("A", opt_a),
                    ("B", opt_b),
                    ("C", opt_c),
                    ("D", opt_d),
                    ("E", opt_e),
                    ("F", opt_f)
                ]
                filtered_opts = [(k, v) for k, v in raw_opts if v is not None]

                if len(filtered_opts) < 2:
                    errors.append(BulkUploadValidationErrorItem(
                        row=row_num,
                        field="option_a",
                        message="MCQ Single Choice questions require at least 2 options (option_a and option_b)."
                    ))

                if not correct_ans_raw:
                    errors.append(BulkUploadValidationErrorItem(
                        row=row_num,
                        field="correct_answer",
                        message="Correct answer is required for MCQ Single Choice questions."
                    ))
                else:
                    ans_clean = correct_ans_raw.strip()
                    if "," in ans_clean:
                        errors.append(BulkUploadValidationErrorItem(
                            row=row_num,
                            field="correct_answer",
                            message=f"MCQ Single Choice questions must have exactly 1 correct answer (found multiple: '{ans_clean}')."
                        ))
                    else:
                        ans_upper = ans_clean.upper()
                        matched_key = None
                        for idx_opt, (k, opt_txt) in enumerate(filtered_opts, start=1):
                            if ans_upper == k or ans_clean == str(idx_opt) or ans_clean.lower() == opt_txt.lower():
                                matched_key = k
                                break

                        if not matched_key:
                            available_keys = [k for k, _ in filtered_opts]
                            errors.append(BulkUploadValidationErrorItem(
                                row=row_num,
                                field="correct_answer",
                                message=f"Correct answer '{ans_clean}' does not match any available option ({', '.join(available_keys)})."
                            ))
                        else:
                            for idx_opt, (k, opt_txt) in enumerate(filtered_opts, start=1):
                                is_corr = (k == matched_key)
                                options_list.append(QuestionOptionCreate(
                                    option_text=opt_txt,
                                    option_order=idx_opt,
                                    option_is_correct=is_corr
                                ))

            elif q_type == QuestionType.MCQ_MULTIPLE:
                opt_a = get_cell("option_a")
                opt_b = get_cell("option_b")
                opt_c = get_cell("option_c")
                opt_d = get_cell("option_d")
                opt_e = get_cell("option_e")
                opt_f = get_cell("option_f")

                raw_opts = [
                    ("A", opt_a),
                    ("B", opt_b),
                    ("C", opt_c),
                    ("D", opt_d),
                    ("E", opt_e),
                    ("F", opt_f)
                ]
                filtered_opts = [(k, v) for k, v in raw_opts if v is not None]

                if len(filtered_opts) < 2:
                    errors.append(BulkUploadValidationErrorItem(
                        row=row_num,
                        field="option_a",
                        message="MCQ Multiple Choice questions require at least 2 options (option_a and option_b)."
                    ))

                if not correct_ans_raw:
                    errors.append(BulkUploadValidationErrorItem(
                        row=row_num,
                        field="correct_answer",
                        message="At least one correct answer is required for MCQ Multiple Choice questions."
                    ))
                else:
                    ans_tokens = [tok.strip() for tok in re.split(r"[,;]+", correct_ans_raw) if tok.strip()]
                    if not ans_tokens:
                        errors.append(BulkUploadValidationErrorItem(
                            row=row_num,
                            field="correct_answer",
                            message="At least one correct answer must be specified."
                        ))
                    else:
                        matched_keys = set()
                        invalid_tokens = []
                        for tok in ans_tokens:
                            tok_upper = tok.upper()
                            found_k = None
                            for idx_opt, (k, opt_txt) in enumerate(filtered_opts, start=1):
                                if tok_upper == k or tok == str(idx_opt) or tok.lower() == opt_txt.lower():
                                    found_k = k
                                    break
                            if found_k:
                                matched_keys.add(found_k)
                            else:
                                invalid_tokens.append(tok)

                        if invalid_tokens:
                            available_keys = [k for k, _ in filtered_opts]
                            errors.append(BulkUploadValidationErrorItem(
                                row=row_num,
                                field="correct_answer",
                                message=f"Answer '{', '.join(invalid_tokens)}' does not match any available option ({', '.join(available_keys)})."
                            ))
                        elif len(matched_keys) < 1:
                            errors.append(BulkUploadValidationErrorItem(
                                row=row_num,
                                field="correct_answer",
                                message="MCQ Multiple Choice questions must have at least 1 correct option."
                            ))
                        else:
                            for idx_opt, (k, opt_txt) in enumerate(filtered_opts, start=1):
                                is_corr = (k in matched_keys)
                                options_list.append(QuestionOptionCreate(
                                    option_text=opt_txt,
                                    option_order=idx_opt,
                                    option_is_correct=is_corr
                                ))

            elif q_type == QuestionType.TRUE_FALSE:
                if not correct_ans_raw:
                    errors.append(BulkUploadValidationErrorItem(
                        row=row_num,
                        field="correct_answer",
                        message="Correct answer ('TRUE' or 'FALSE') is required for True/False questions."
                    ))
                else:
                    ans_norm = correct_ans_raw.strip().upper()
                    if ans_norm in ["TRUE", "T", "1", "YES"]:
                        options_list = [
                            QuestionOptionCreate(option_text="True", option_order=1, option_is_correct=True),
                            QuestionOptionCreate(option_text="False", option_order=2, option_is_correct=False),
                        ]
                    elif ans_norm in ["FALSE", "F", "0", "NO"]:
                        options_list = [
                            QuestionOptionCreate(option_text="True", option_order=1, option_is_correct=False),
                            QuestionOptionCreate(option_text="False", option_order=2, option_is_correct=True),
                        ]
                    else:
                        errors.append(BulkUploadValidationErrorItem(
                            row=row_num,
                            field="correct_answer",
                            message=f"Invalid True/False answer '{correct_ans_raw}'. Must be 'TRUE' or 'FALSE'."
                        ))

            elif q_type == QuestionType.SHORT_ANSWER:
                options_list = []

            # If all validations for this row passed, add to parsed list
            row_has_errors = any(e.row == row_num for e in errors)
            if not row_has_errors and q_text and q_type and marks_val is not None:
                try:
                    validate_question_options_rules(q_type, options_list)
                except ValidationException as ve:
                    errors.append(BulkUploadValidationErrorItem(
                        row=row_num,
                        field="options",
                        message=str(ve.detail if hasattr(ve, "detail") else ve)
                    ))
                    continue

                total_marks += marks_val
                parsed_questions.append({
                    "row": row_num,
                    "question_text": q_text,
                    "question_type": q_type,
                    "marks": marks_val,
                    "explanation": explanation,
                    "order": order_val,
                    "options": options_list,
                    "correct_answer": correct_ans_raw or ("True" if q_type == QuestionType.TRUE_FALSE and options_list and options_list[0].option_is_correct else "False" if q_type == QuestionType.TRUE_FALSE else None)
                })

        return parsed_questions, errors, total_marks
